"""Upload → map → validate → commit, with original and processed downloads."""

import csv
import hashlib
import io
import json
import re
from datetime import UTC, date, datetime
from pathlib import PurePath
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile, ZipFile

from defusedxml.common import DefusedXmlException
from fastapi import APIRouter, Depends, HTTPException, Response, UploadFile
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from pydantic import Field, ValidationError
from sqlalchemy.orm import Session

from app.auth import get_current_user
from app.db import get_db
from app.models import Contract, ImportBatch
from app.routers.domain import (
    EvidenceIn,
    InvoiceIn,
    LineIn,
    StrictModel,
    add_evidence,
    add_invoice,
    contract_for,
)
from app.services.access import audit, entity_scope, fields, scope

router = APIRouter(prefix="/ingestion", tags=["file processing"])
MAX_BYTES = 2 * 1024 * 1024
MAX_ROWS = 500
BASE_COLUMNS = ("work_reference", "work_code", "quantity", "service_date", "unit")


def required(kind):
    return (
        (("supplier_invoice_no",) + BASE_COLUMNS + ("unit_price",))
        if kind == "invoice"
        else BASE_COLUMNS
    )


def roles(kind):
    return (
        ("ap_analyst", "tenant_admin")
        if kind == "invoice"
        else ("operations_verifier", "tenant_admin")
    )


def cell_text(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return "" if value is None else str(value).strip()


def extract_file(raw, filename):
    suffix = PurePath(filename).suffix.lower()
    sheet = None
    if suffix == ".csv":
        try:
            text = raw.decode("utf-8-sig")
            if "\x00" in text:
                raise ValueError("Binary content is not accepted as CSV")
            values = list(csv.reader(io.StringIO(text), strict=True))
        except (UnicodeError, csv.Error, ValueError) as exc:
            raise HTTPException(422, "Invalid UTF-8 CSV: " + str(exc)) from exc
    elif suffix == ".xlsx":
        try:
            with ZipFile(io.BytesIO(raw)) as archive:
                entries = archive.infolist()
                if (
                    len(entries) > 1000
                    or sum(e.file_size for e in entries) > 16 * 1024 * 1024
                ):
                    raise ValueError(
                        "Workbook expanded size exceeds the 16 MiB processing limit"
                    )
                if any(
                    "vbaProject" in e.filename or "externalLinks/" in e.filename
                    for e in entries
                ):
                    raise ValueError(
                        "Macros and external workbook links are not supported"
                    )
            book = load_workbook(
                io.BytesIO(raw), read_only=True, data_only=False, keep_links=False
            )
            try:
                worksheet = book.worksheets[0]
                sheet = worksheet.title
                values = []
                for row in worksheet.iter_rows():
                    if len(values) > MAX_ROWS or len(row) > 32:
                        raise ValueError(
                            "Workbook limit is 500 data rows and 32 columns"
                        )
                    if any(cell.data_type == "f" for cell in row):
                        raise ValueError(
                            "Formula cells must be replaced with reviewed values before ingestion"
                        )
                    values.append([cell_text(cell.value) for cell in row])
            finally:
                book.close()
        except (
            BadZipFile,
            ValueError,
            KeyError,
            IndexError,
            OSError,
            ParseError,
            InvalidFileException,
            DefusedXmlException,
        ) as exc:
            raise HTTPException(422, "Invalid XLSX: " + str(exc)) from exc
    else:
        raise HTTPException(422, "Upload a UTF-8 CSV or an .xlsx workbook")
    indexed_values = [
        (number, row)
        for number, row in enumerate(values, 1)
        if any(str(v).strip() for v in row)
    ]
    if len(indexed_values) < 2:
        raise HTTPException(422, "File must contain a header and at least one data row")
    if len(indexed_values) > MAX_ROWS + 1:
        raise HTTPException(413, "File limit is 500 data rows")
    headers = [cell_text(v) for v in indexed_values[0][1]]
    if (
        len(headers) > 32
        or any(not h for h in headers)
        or len(set(headers)) != len(headers)
    ):
        raise HTTPException(422, "Use 1–32 unique, nonempty column headers")
    return {
        "headers": headers,
        "raw_rows": [[cell_text(v) for v in row] for _, row in indexed_values[1:]],
        "source_rows": [number for number, _ in indexed_values[1:]],
        "sheet": sheet,
        "rows": [],
    }


def batch_detail(db, batch):
    meta = json.loads(batch.errors_json or "{}")
    data = json.loads(batch.normalized_json or "{}")
    return {
        **fields(
            batch,
            "id",
            "org_id",
            "site_id",
            "kind",
            "filename",
            "sha256",
            "status",
            "created_at",
        ),
        "contract_id": meta.get("contract_id"),
        "accepted": meta.get("accepted", 0),
        "errors": meta.get("errors", []),
        "created_records": meta.get("created_records", []),
        "headers": data.get("headers", []),
        "required_columns": required(batch.kind),
        "mapping": json.loads(batch.mapping_json or "{}"),
        "sheet": data.get("sheet"),
        "row_count": len(data.get("raw_rows", [])),
        "preview": data.get("raw_rows", [])[:10],
        "normalized_preview": data.get("rows", [])[:10],
        "contract": fields(db.get(Contract, meta["contract_id"]), "id", "title")
        if meta.get("contract_id") and db.get(Contract, meta["contract_id"])
        else None,
        "source_retained": batch.purged_at is None,
        "purged_at": batch.purged_at.isoformat() if batch.purged_at else None,
        "purged_by": batch.purged_by,
    }


def validate_batch(batch, mapping):
    data = json.loads(batch.normalized_json)
    meta = json.loads(batch.errors_json)
    columns = required(batch.kind)
    errors, normalized = [], []
    if set(mapping) != set(columns) or any(
        mapping[c] not in data["headers"] for c in columns
    ):
        raise HTTPException(422, "Map every required field to a source column")
    if len(set(mapping.values())) != len(mapping):
        raise HTTPException(
            422, "A source column cannot map to multiple required fields"
        )
    source_rows = data.get("source_rows", list(range(2, len(data["raw_rows"]) + 2)))
    for number, values in zip(source_rows, data["raw_rows"], strict=True):
        if len(values) != len(data["headers"]):
            errors.append(
                {
                    "row": number,
                    "field": "row",
                    "message": "Column count differs from header",
                }
            )
            continue
        source = dict(zip(data["headers"], values, strict=True))
        row = {column: source[mapping[column]] for column in columns}
        try:
            if row["unit"] != "hour":
                raise ValueError("Only hour units are supported")
            if batch.kind == "invoice":
                reference = row.pop("supplier_invoice_no").strip()
                if not reference or len(reference) > 200:
                    raise ValueError("Supplier invoice number must be 1–200 characters")
                parsed = LineIn.model_validate(row).model_dump(mode="json")
                parsed["supplier_invoice_no"] = reference
            else:
                parsed = EvidenceIn.model_validate(
                    {
                        "org_id": batch.org_id,
                        "site_id": batch.site_id,
                        "contract_id": meta["contract_id"],
                        **row,
                    }
                ).model_dump(mode="json")
            normalized.append({"source_row": number, **parsed})
        except ValidationError as exc:
            errors.extend(
                {
                    "row": number,
                    "field": ".".join(map(str, e["loc"])),
                    "message": e["msg"],
                }
                for e in exc.errors()
            )
        except ValueError as exc:
            errors.append({"row": number, "field": "row", "message": str(exc)})
    data["rows"] = normalized
    meta.update(errors=errors, accepted=0)
    batch.status = "rejected" if errors else "ready"
    batch.mapping_json = json.dumps(mapping)
    batch.normalized_json = json.dumps(data)
    batch.errors_json = json.dumps(meta)


@router.get("")
def files(
    org_id: str,
    site_id: str,
    user=Depends(get_current_user),
    db: Session = Depends(get_db),
):
    scope(db, user, org_id, site_id)
    return [
        batch_detail(db, batch)
        for batch in db.query(ImportBatch)
        .filter_by(org_id=org_id, site_id=site_id)
        .order_by(ImportBatch.created_at.desc())
        .limit(200)
    ]


@router.post("/{kind}/stage")
def stage_file(
    kind: str,
    org_id: str,
    site_id: str,
    contract_id: str,
    file: UploadFile,
    user=Depends(get_current_user),
    db: Session = Depends(get_db),
):
    if kind not in ("invoice", "evidence"):
        raise HTTPException(422, "Choose invoice or evidence ingestion")
    contract_for(db, user, contract_id, org_id, site_id, *roles(kind))
    db.query(Contract).filter_by(id=contract_id).with_for_update().one()
    raw = file.file.read(MAX_BYTES + 1)
    if len(raw) > MAX_BYTES:
        raise HTTPException(413, "File limit is 2 MiB")
    filename = (file.filename or "upload")[:200]
    data = extract_file(raw, filename)
    digest = hashlib.sha256(raw).hexdigest()
    for prior in db.query(ImportBatch).filter_by(
        org_id=org_id, site_id=site_id, kind=kind, sha256=digest
    ):
        if (
            prior.normalized_json
            and json.loads(prior.errors_json or "{}").get("contract_id") == contract_id
        ):
            return batch_detail(db, prior)
    batch = ImportBatch(
        org_id=org_id,
        site_id=site_id,
        kind=kind,
        filename=filename,
        sha256=digest,
        source_bytes=raw,
        source_text=raw.decode("utf-8-sig")
        if filename.lower().endswith(".csv")
        else None,
        created_by=user.id,
        status="staged",
        normalized_json=json.dumps(data),
        errors_json=json.dumps(
            {"contract_id": contract_id, "errors": [], "accepted": 0}
        ),
    )
    db.add(batch)
    db.flush()
    if all(column in data["headers"] for column in required(kind)):
        validate_batch(batch, {column: column for column in required(kind)})
    audit(
        db,
        user,
        batch,
        "file.staged",
        {"sha256": digest, "rows": len(data["raw_rows"])},
    )
    db.commit()
    return batch_detail(db, batch)


@router.get("/{batch_id}")
def get_batch(
    batch_id: str, user=Depends(get_current_user), db: Session = Depends(get_db)
):
    return batch_detail(db, entity_scope(db, user, db.get(ImportBatch, batch_id)))


class MappingIn(StrictModel):
    mapping: dict[str, str] = Field(max_length=10)


@router.post("/{batch_id}/validate")
def validate_file(
    batch_id: str,
    body: MappingIn,
    user=Depends(get_current_user),
    db: Session = Depends(get_db),
):
    batch = db.query(ImportBatch).filter_by(id=batch_id).with_for_update().first()
    entity_scope(db, user, batch, *roles(batch.kind) if batch else ())
    if batch.status == "imported" or not batch.normalized_json:
        raise HTTPException(409, "Imported or legacy files cannot be remapped")
    validate_batch(batch, body.mapping)
    audit(db, user, batch, "file.validated", {"status": batch.status})
    db.commit()
    return batch_detail(db, batch)


@router.post("/{batch_id}/commit")
def commit_file(
    batch_id: str, user=Depends(get_current_user), db: Session = Depends(get_db)
):
    batch = db.query(ImportBatch).filter_by(id=batch_id).with_for_update().first()
    entity_scope(db, user, batch, *roles(batch.kind) if batch else ())
    if batch.status == "imported":
        return batch_detail(db, batch)
    if batch.status != "ready":
        raise HTTPException(409, "Validate all rows successfully before importing")
    meta = json.loads(batch.errors_json)
    contract_for(
        db, user, meta["contract_id"], batch.org_id, batch.site_id, *roles(batch.kind)
    )
    data = json.loads(batch.normalized_json)
    created = []
    try:
        with db.begin_nested():
            if batch.kind == "evidence":
                for row in data["rows"]:
                    record = add_evidence(
                        db,
                        user,
                        EvidenceIn.model_validate(
                            {k: v for k, v in row.items() if k != "source_row"}
                        ),
                    )
                    record.import_batch_id = batch.id
                    created.append(
                        {
                            "kind": "evidence",
                            "id": record.id,
                            "label": record.work_reference,
                        }
                    )
            else:
                grouped = {}
                for row in data["rows"]:
                    grouped.setdefault(row["supplier_invoice_no"], []).append(
                        LineIn.model_validate(
                            {
                                k: v
                                for k, v in row.items()
                                if k not in ("source_row", "supplier_invoice_no")
                            }
                        )
                    )
                for reference, lines in grouped.items():
                    record = add_invoice(
                        db,
                        user,
                        InvoiceIn(
                            org_id=batch.org_id,
                            site_id=batch.site_id,
                            contract_id=meta["contract_id"],
                            supplier_invoice_no=reference,
                            lines=lines,
                        ),
                    )
                    record.import_batch_id = batch.id
                    created.append(
                        {"kind": "invoices", "id": record.id, "label": reference}
                    )
    except HTTPException as exc:
        batch.status = "rejected"
        meta.update(
            errors=[{"row": 0, "field": "domain", "message": str(exc.detail)}],
            accepted=0,
            created_records=[],
        )
    else:
        batch.status = "imported"
        meta.update(errors=[], accepted=len(data["rows"]), created_records=created)
    batch.errors_json = json.dumps(meta)
    audit(db, user, batch, "file." + batch.status, {"accepted": meta["accepted"]})
    db.commit()
    return batch_detail(db, batch)


@router.post("/{batch_id}/purge")
def purge_source(
    batch_id: str, user=Depends(get_current_user), db: Session = Depends(get_db)
):
    """Remove original uploaded bytes after import; keep checksum and lineage."""
    batch = db.query(ImportBatch).filter_by(id=batch_id).with_for_update().first()
    entity_scope(db, user, batch, "tenant_admin")
    if batch.purged_at is not None:
        return batch_detail(db, batch)
    if batch.status != "imported":
        raise HTTPException(
            409, "Only imported batches can be purged; commit or discard the batch"
        )
    size = len(batch.source_bytes or b"") or len(
        (batch.source_text or "").encode("utf-8")
    )
    batch.source_bytes = None
    batch.source_text = None
    batch.purged_at = datetime.now(UTC)
    batch.purged_by = user.id
    audit(
        db,
        user,
        batch,
        "file.source_purged",
        {"sha256": batch.sha256, "bytes": size},
    )
    db.commit()
    return batch_detail(db, batch)


def csv_bytes(headers, rows):
    stream = io.StringIO(newline="")
    writer = csv.writer(stream)
    writer.writerow(headers)
    for row in rows:
        # Downloads open safely as data in spreadsheet applications.
        writer.writerow(
            [
                "'" + str(v) if str(v).startswith(("=", "+", "-", "@")) else v
                for v in row
            ]
        )
    return stream.getvalue().encode("utf-8-sig")


@router.get("/templates/{kind}")
def template_download(kind: str, user=Depends(get_current_user)):
    if kind not in ("evidence", "invoice"):
        raise HTTPException(422, "Choose evidence or invoice template")
    sample = {
        "supplier_invoice_no": "EXAMPLE-INVOICE",
        "work_reference": "EXAMPLE-SHIFT",
        "work_code": "EQ-HIRE",
        "quantity": "160" if kind == "invoice" else "145",
        "service_date": "2026-10-10",
        "unit": "hour",
        "unit_price": "250",
    }
    columns = required(kind)
    return Response(
        csv_bytes(columns, [[sample[column] for column in columns]]),
        media_type="text/csv",
        headers={"Content-Disposition": f'attachment; filename="{kind}-template.csv"'},
    )


@router.get("/{batch_id}/download/{artifact}")
def file_download(
    batch_id: str,
    artifact: str,
    user=Depends(get_current_user),
    db: Session = Depends(get_db),
):
    batch = entity_scope(db, user, db.get(ImportBatch, batch_id))
    meta = json.loads(batch.errors_json or "{}")
    data = json.loads(batch.normalized_json or "{}")
    if artifact == "source":
        if batch.purged_at is not None:
            raise HTTPException(
                410,
                "Original file content was purged; checksum and processed rows remain",
            )
        content = batch.source_bytes or (batch.source_text or "").encode()
        filename = re.sub(r"[^A-Za-z0-9_.-]", "_", batch.filename)
        mime = "application/octet-stream"
    elif artifact == "errors":
        content = csv_bytes(
            ["row", "field", "message"],
            [
                [e.get("row"), e.get("field", "row"), e["message"]]
                for e in meta.get("errors", [])
            ],
        )
        filename, mime = "validation-errors.csv", "text/csv"
    elif artifact == "normalized":
        columns = required(batch.kind)
        content = csv_bytes(
            ["source_row", *columns],
            [
                [r["source_row"], *[r.get(c, "") for c in columns]]
                for r in data.get("rows", [])
            ],
        )
        filename, mime = "processed-rows.csv", "text/csv"
    else:
        raise HTTPException(404, "Choose source, normalized or errors output")
    return Response(
        content,
        media_type=mime,
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
